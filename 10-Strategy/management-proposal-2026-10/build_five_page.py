#!/usr/bin/env python3
"""Generate the concise Marindustrial management brief from its tracked Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DIR = Path(__file__).resolve().parent
SRC = DIR / "Marindustrial_5_Page_Management_Brief.md"
DEST = DIR / "Marindustrial_5_Page_Management_Brief.docx"
NAVY = "193544"
TEAL = "118B92"
DARK = "1F3640"
MUTED = "526670"
PALE = "EAF3F5"
LIGHT = "F3F7F8"
WHITE = "FFFFFF"

def rgb(s): return RGBColor.from_string(s)

def shade(cell, fill):
    e = OxmlElement("w:shd")
    e.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(e)

def pad(cell, top=140, start=170, bottom=140, end=170):
    tc = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, n in (("top",top),("start",start),("bottom",bottom),("end",end)):
        e = OxmlElement("w:"+side)
        e.set(qn("w:w"), str(n))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tc.append(mar)

def norules(tab):
    e = OxmlElement("w:tblBorders")
    for side in ("top","bottom","left","right","insideV","insideH"):
        s = OxmlElement("w:"+side)
        s.set(qn("w:val"), "nil")
        e.append(s)
    tab._tbl.tblPr.append(e)

def text_run(p, content, size=10, color=DARK, bold=False, italic=False):
    r = p.add_run(content)
    r.bold = bold
    r.italic = italic
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = rgb(color)
    return r

def paragraph(parent, content, size=10, after=7, color=DARK, bold=False, before=0, italic=False):
    p = parent.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.14
    text_run(p, content, size, color, bold, italic)
    return p

def rich(parent, line, size=9.4, after=5):
    p = parent.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.17
    chunks = re.split(r"(\*\*[^*]+\*\*)", line)
    for part in chunks:
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            text_run(p, part[2:-2], size, NAVY, True)
        else: text_run(p, part, size, DARK)
    return p

def callout(doc, line):
    tab = doc.add_table(rows=1, cols=1)
    c = tab.cell(0, 0)
    shade(c, NAVY)
    pad(c, 180, 230, 160, 230)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    for j,part in enumerate(re.split(r"(\*\*[^*]+\*\*)", line.removeprefix("> ").strip())):
        if not part: continue
        if part.startswith("**"):
            text_run(p, part[2:-2], 9.25, "A8E0E3", True)
        else:
            text_run(p, part, 9.55, WHITE)
    norules(tab)

def header(doc, page, title, subtitle):
    top = paragraph(doc, f"MANAGEMENT BRIEF  /  OCTOBER 2026    •    {page:02}/05",
                    size=8, color=TEAL, bold=True, after=12)
    paragraph(doc, title, size=23.2 if page != 1 else 26, color=NAVY, bold=True, after=10)
    if subtitle:
        paragraph(doc, subtitle, size=10.3, color=MUTED, italic=True, after=15)
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0,0)
    shade(c, TEAL)
    pad(c,0,0,0,0)
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    t.rows[0].height = Inches(.028)
    t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    norules(t)
    paragraph(doc, "", size=1, after=9)

def make_card(doc, title, fields):
    match = re.match(r"(\d\d) \| (.*)",title)
    num,name = match.groups()
    t = doc.add_table(rows=1,cols=2)
    t.autofit = False
    t.columns[0].width = Inches(.56)
    t.columns[1].width = Inches(6.35)
    l,r = t.rows[0].cells
    shade(l,NAVY)
    shade(r,LIGHT)
    pad(l,145,85,115,60)
    pad(r,145,180,125,185)
    l.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    text_run(l.paragraphs[0],num,11,WHITE,True)
    text_run(r.paragraphs[0],name,12.0,NAVY,True)
    r.paragraphs[0].paragraph_format.space_after = Pt(8)
    for line in fields:
        p = rich(r, line, 9.05, 4)
        # The initial empty paragraph in a cell is never duplicated.
    norules(t)
    paragraph(doc, "", size=1, after=8)

lines = SRC.read_text(encoding="utf-8").splitlines()
pages = "\n".join(lines).split("[[PAGE_BREAK]]")
if len(pages) != 5: raise RuntimeError(f"Expected 5 source sections; found {len(pages)}")

doc = Document()
sec = doc.sections[0]
sec.page_height = Inches(11)
sec.page_width = Inches(8.5)
sec.left_margin = Inches(.76)
sec.right_margin = Inches(.74)
sec.top_margin = Inches(.76)
sec.bottom_margin = Inches(.67)
sec.header_distance = Inches(.32)
sec.footer_distance = Inches(.36)
normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9.7)
normal.font.color.rgb = rgb(DARK)
normal.paragraph_format.space_after = Pt(0)
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
text_run(hp,"MARINDUSTRIAL  /  OPERATIONS IMPROVEMENT",7.4,MUTED,True)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
text_run(fp,"MANAGEMENT REVIEW  •  OCT 2026     |     PAGE ",7.2,MUTED)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"),"PAGE")
fp._p.append(fld)

for ix, source in enumerate(pages, start=1):
    if ix>1: doc.add_page_break()
    raw = [x.rstrip() for x in source.splitlines() if x.strip()]
    first = next((a[2:] for a in raw if a.startswith("# ")), "Management Proposal")
    subtitles = [x.strip("*") for x in raw if x.startswith("*") and not x.startswith("**")]
    header(doc, ix, first, subtitles[0] if subtitles else "")
    i=0
    while i<len(raw):
        l=raw[i]
        if l.startswith("# ") or (l.startswith("*") and not l.startswith("**")) or l=="---":
            i+=1;continue
        if re.match(r"## \d\d \|",l):
            j=i+1
            fields=[]
            while j<len(raw) and not raw[j].startswith(("## ","# ","[[PAGE_BREAK]]")):
                if raw[j].startswith("**"):fields.append(raw[j])
                j+=1
            make_card(doc,l[3:],fields)
            i=j;continue
        if l.startswith("## "):
            paragraph(doc, l[3:].upper(), 9.55, 8, TEAL, True, before=9)
        elif l.startswith("### "):
            paragraph(doc, l[4:], 10.2, 6, NAVY, True, before=5)
        elif l.startswith("> "):
            callout(doc,l)
            paragraph(doc,"",size=1,after=3)
        elif re.match(r"\d+\. ",l) or l.startswith("- "):
            rich(doc, "•  " + re.sub(r"^(?:\d+\.|-)\s+", "",l), size=9.45, after=6)
        else:
            rich(doc,l, size=9.65 if ix==1 else 9.4, after=9)
        i+=1

props = doc.core_properties
props.title = "Marindustrial | Five-page Management Brief"
props.subject = "Controlled internal software development evaluation"
props.author = "Marindustrial proposal working draft"
doc.save(DEST)
print(DEST)
